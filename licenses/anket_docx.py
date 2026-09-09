"""
Lisenziya/İdxal-İxrac və s. sənəd yaratma formalarındakı "2. Anket" bölməsi üçün
Word (.docx) şablon generasiyası və doldurulmuş şablonun geri oxunması.

Axın:
  1. İstifadəçi "Şablonu yüklə" düyməsini basır -> build_anket_template() bir .docx
     yaradır: cədvəldə "Sahə adı | Dəyər" sütunları var. VÖEN və Müəssisənin adı
     artıq API-dən bilindiyi üçün əvvəlcədən doldurulur, "auto"/"readonly" sahələr
     də (məs. "Müddət") sabit dəyərləri ilə doldurulur - istifadəçi yalnız qalan
     boş sətirləri doldurur.
  2. İstifadəçi Word-də doldurub geri yükləyir -> parse_anket_template() cədvəli
     oxuyub sətir başlıqlarını (Sahə adı) schema-dakı form_fields-in "label"/"key"
     dəyərləri ilə uyğunlaşdırır və {field_key: deyer} formasında qaytarır.

QEYD: Uyğunlaşdırma mətn əsaslıdır (normallaşdırılmış müqayisə) - Word sənədində
sətir sırası və ya əlavə boşluqlar problem yaratmır, AMMA sətrin birinci sütunu
(Sahə adı) mümkün qədər orijinal etiketə (və ya field key-inə) yaxın olmalıdır.

"checkbox_list" tipli sahələr (məs. "İstinad maddəsi" bəndləri) üçün hər bir
seçim (bənd) öz sətrində "☐"/"☑" işarəsi ilə göstərilir - istifadəçi aid olan
bəndin qarşısındakı "☐" işarəsini "☑" və ya "X" ilə əvəz edir. Bu, real Word
"content control" checkbox-larından fərqli olaraq sadə mətndir, ona görə bütün
Word versiyalarında (o cümlədən mobil/onlayn) problemsiz açılır və geri
yükləndikdə etibarlı şəkildə oxuna bilir.

"computed" (sistem tərəfindən avtomatik hesablanan, məs. "Lisenziya
kateqoriyası") sahələr şablonda GÖSTƏRİLMİR - onların dəyəri yalnız sənəd
göndəriləndən sonra bəndlərə əsasən müəyyən olunur.
"""

import io
import re

from docx import Document

_CHECKED_MARKS = ("☑", "✓", "✔", "x", "X", "V", "v", "+")


def _normalize_label(text):
    """Azərbaycan hərflərini latın uyğunları ilə əvəz edib, boşluq/durğu işarələrini
    tək alt xəttə çevirərək müqayisə üçün "normal" mətn açarı yaradır."""
    text = (text or "").strip().lower()
    translit = str.maketrans({
        "ə": "e", "ö": "o", "ü": "u", "ç": "c", "ş": "s", "ğ": "g", "ı": "i",
    })
    text = text.translate(translit)
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def build_anket_template(schema, applicant_name="", voen="", title="Lisenziya anketi"):
    """schema: get_schema(doc_type) nəticəsi (dict, 'form_fields' açarı ilə).
    Qaytarır: BytesIO (.docx məzmunu)."""

    document = Document()

    document.add_heading(title, level=1)

    intro = document.add_paragraph(
        "Bu sənəd sistem tərəfindən avtomatik yaradılıb. Aşağıdakı cədvəldə "
        "\"Dəyər\" sütununda boş qalan sətirləri doldurub sənədi olduğu formatda "
        "(Word/.docx) geri yükləyin - sistem dəyərləri avtomatik anketə köçürəcək."
    )
    intro.italic = True

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    header_cells[0].text = "Sahə adı"
    header_cells[1].text = "Dəyər"

    # VÖEN və müəssisə adı - "Müraciətçi məlumatları" bölümündən, API-dən artıq bəllidir.
    voen_row = table.add_row().cells
    voen_row[0].text = "VÖEN"
    voen_row[1].text = voen or ""

    if applicant_name:
        name_row = table.add_row().cells
        name_row[0].text = "Müəssisənin adı"
        name_row[1].text = applicant_name

    for field in schema.get("form_fields", []):
        # Sistem tərəfindən avtomatik hesablanan sahələr (məs. "Lisenziya
        # kateqoriyası") şablonda görünmür - istifadəçi bunları doldurmur.
        if field.get("computed"):
            continue

        if field.get("type") == "checkbox_list":
            for opt in field.get("options", []):
                opt_key = opt["key"] if isinstance(opt, dict) else opt[0]
                opt_label = opt["label"] if isinstance(opt, dict) else opt[1]
                row = table.add_row().cells
                row[0].text = f"{field.get('label', field.get('key', ''))} — {opt_label}"
                row[1].text = "☐"
            continue

        row = table.add_row().cells
        label = field.get("label", field.get("key", ""))
        options = field.get("options")
        if options:
            # "select" sahələr üçün Word-də hansı dəyərlərin qəbul olunduğunu göstəririk
            # ki, istifadəçi sərbəst mətn yox, məhz bu variantlardan birini yazsın.
            option_labels = " / ".join(opt_label for _, opt_label in options)
            label = f"{label} (seçim: {option_labels})"
        row[0].text = label
        # "auto"/"readonly" sahələrin sabit dəyəri varsa, əvvəlcədən doldururuq -
        # istifadəçi bunları yenidən yazmasın deyə.
        fixed_value = field.get("value")
        row[1].text = str(fixed_value) if fixed_value not in (None, "") else ""

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def parse_anket_template(file_obj, schema):
    """file_obj: yüklənmiş .docx faylı (Django UploadedFile və ya BytesIO).
    schema: get_schema(doc_type) nəticəsi.

    Qaytarır: {"values": {field_key: deyer, ...}, "voen": "..." | None,
    "matched_count": N}."""

    document = Document(file_obj)

    if not document.tables:
        return {"values": {}, "voen": None, "matched_count": 0}

    table = document.tables[0]

    # Sahə etiketlərini (label/key) normallaşdırıb axtarış üçün lüğət qururuq.
    field_lookup = {}
    field_by_key = {}
    # checkbox_list sahələr üçün ayrıca lüğət: "sahə etiketi — bənd etiketi" (normallaşmış) -> (field_key, option_key)
    checkbox_lookup = {}
    checkbox_field_keys = set()
    for field in schema.get("form_fields", []):
        if field.get("computed"):
            continue  # sistem hesablayır, şablonda sətri yoxdur
        field_by_key[field["key"]] = field
        if field.get("type") == "checkbox_list":
            checkbox_field_keys.add(field["key"])
            field_label = field.get("label", field.get("key", ""))
            for opt in field.get("options", []):
                opt_key = opt["key"] if isinstance(opt, dict) else opt[0]
                opt_label = opt["label"] if isinstance(opt, dict) else opt[1]
                norm = _normalize_label(f"{field_label} — {opt_label}")
                if norm:
                    checkbox_lookup[norm] = (field["key"], opt_key)
            continue
        for candidate in (field.get("label"), field.get("key")):
            norm = _normalize_label(candidate)
            if norm:
                field_lookup[norm] = field["key"]

    values = {}
    checked_bends = {key: [] for key in checkbox_field_keys}
    voen_value = None

    for row in table.rows[1:]:  # ilk sətir header - keçirik
        cells = row.cells
        if len(cells) < 2:
            continue

        raw_label = cells[0].text.strip()
        raw_value = cells[1].text.strip()

        if not raw_label or not raw_value:
            continue

        # Template generasiyasında "select" sahələrin etiketinə " (seçim: ...)"
        # əlavə olunur (istifadəçi üçün ipucu) - uyğunlaşdırmadan əvvəl bunu ataq.
        raw_label = re.sub(r"\s*\(seçim:.*?\)\s*$", "", raw_label, flags=re.IGNORECASE)

        norm_label = _normalize_label(raw_label)

        if norm_label == "voen":
            voen_value = raw_value
            continue

        if norm_label in ("muessisenin_adi", "musessisenin_adi", "muraciet_edenin_adi"):
            continue  # applicant_name ayrıca idarə olunur, anket sahəsi deyil

        # checkbox_list sətri olub-olmadığını yoxlayırıq (məs. "İstinad maddəsi — VI bənd").
        checkbox_match = checkbox_lookup.get(norm_label)
        if not checkbox_match:
            for norm_candidate, pair in checkbox_lookup.items():
                if norm_candidate and (norm_candidate in norm_label or norm_label in norm_candidate):
                    checkbox_match = pair
                    break
        if checkbox_match:
            field_key, option_key = checkbox_match
            is_checked = any(mark in raw_value for mark in _CHECKED_MARKS)
            if is_checked:
                checked_bends[field_key].append(option_key)
            continue

        matched_key = field_lookup.get(norm_label)

        if not matched_key:
            # Tam uyğunluq tapılmadısa, qismən uyğunluğa (substring) cəhd edirik.
            for norm_candidate, key in field_lookup.items():
                if norm_candidate and (norm_candidate in norm_label or norm_label in norm_candidate):
                    matched_key = key
                    break

        if matched_key:
            field = field_by_key.get(matched_key)
            options = field.get("options") if field else None
            if options:
                # "select" sahə - istifadəçinin yazdığı mətni (label) option key-inə
                # çeviririk (məs. "Silah" -> "silah"), tapılmasa xam mətni saxlayırıq.
                norm_value = _normalize_label(raw_value)
                option_key = next(
                    (opt_key for opt_key, opt_label in options if _normalize_label(opt_label) == norm_value
                     or opt_key == raw_value),
                    None,
                )
                values[matched_key] = option_key or raw_value
            else:
                values[matched_key] = raw_value

    for field_key, checked_keys in checked_bends.items():
        values[field_key] = checked_keys

    return {"values": values, "voen": voen_value, "matched_count": len(values)}